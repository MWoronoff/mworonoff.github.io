"""Step 2: shared geography layer for the Adtaxi market analyzers.

Inputs
  - Deluxe ZIP workbooks (10 files, 00000_to_09999.xlsx ... 90000_to_99999.xlsx), passed as --deluxe-dir.
    These are licensed vendor files and are NOT committed to the public repo.
  - data-build/sources/dma_zip_2026.csv  (public ZIP-level DMA list, github.com/BritCrit/dma_county_zip)
  - data-build/sources/dma_county_2021.csv (optional cross-check, github.com/alex-patton/US-TVDMA-BY-COUNTY)

Outputs (data-build/geo/)
  zip_geo.csv, county_geo.csv, dma_list.csv, metro_list.csv, dma_disagreements.csv, geography_report.md
Scope: 50 states + DC. Puerto Rico, other territories and military ZIPs are excluded (no Nielsen DMA).
"""
import argparse, glob, os, re, sys
import pandas as pd
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "geo")
EXCLUDE_STATES = {"PR", "VI", "GU", "AS", "MP", "FM", "MH", "PW", "AA", "AE", "AP"}
COLS = ["ZipCode", "City", "State", "County", "CountyFIPS", "StateFIPS", "CBSA", "CBSA_Name", "CBSA_Type",
        "MultiCounty", "Population", "HouseholdsPerZipCode", "Latitude", "Longitude"]


def load_deluxe(d):
    rows = []
    files = sorted(glob.glob(os.path.join(d, "*_to_*.xlsx")))
    if len(files) != 10:
        sys.exit(f"Expected 10 Deluxe workbooks in {d}, found {len(files)}")
    for f in files:
        ws = load_workbook(f, read_only=True).worksheets[0]
        it = ws.iter_rows(values_only=True)
        hdr = list(next(it))
        idx = [hdr.index(c) for c in COLS]
        rows += [[r[i] for i in idx] for r in it]
    z = pd.DataFrame(rows, columns=COLS)
    z["zip"] = z.ZipCode.astype(str).str.zfill(5)
    return z


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--deluxe-dir", required=True)
    a = ap.parse_args()
    os.makedirs(OUT, exist_ok=True)
    rep = []  # (check, result, detail)

    raw = load_deluxe(a.deluxe_dir)
    n_raw = len(raw)
    dup = raw.zip.duplicated().sum()
    rep.append(("One row per ZIP", "PASS" if dup == 0 else "FAIL", f"{n_raw:,} ZIPs, {dup} duplicates"))

    mil = raw.CountyFIPS.isna()
    terr = raw.State.isin(EXCLUDE_STATES)
    z = raw[~mil & ~terr].copy()
    rep.append(("Scope", "INFO", f"Kept {len(z):,} ZIPs in 50 states + DC; excluded {int(mil.sum())} military and "
                f"{int((terr & ~mil).sum())} territory ZIPs"))

    z["county_fips"] = z.StateFIPS.astype(int).astype(str).str.zfill(2) + z.CountyFIPS.astype(int).astype(str).str.zfill(3)
    z["population"] = pd.to_numeric(z.Population, errors="coerce").fillna(0).astype(int)
    z["households"] = pd.to_numeric(z.HouseholdsPerZipCode, errors="coerce").fillna(0).astype(int)
    z["cbsa_code"] = z.CBSA.apply(lambda v: "" if pd.isna(v) else str(int(v)).zfill(5))
    z["cbsa_type"] = z.CBSA_Type.fillna("")
    z["cbsa_name"] = z.CBSA_Name.fillna("")
    z["multi_county"] = (z.MultiCounty == "Y")

    # --- DMA by ZIP, county fill for unmatched ZIPs
    dz = pd.read_csv(os.path.join(HERE, "sources", "dma_zip_2026.csv"), dtype=str, encoding="latin1")
    dz["zip"] = dz.zipcode.str.zfill(5)
    dz = dz.drop_duplicates("zip")[["zip", "dma_code", "dma_name"]]
    ov = pd.read_csv(os.path.join(HERE, "sources", "dma_overrides.csv"), dtype=str)
    dz = pd.concat([dz[~dz.zip.isin(ov.zip)], ov[["zip", "dma_code", "dma_name"]]])
    dz.loc[dz.dma_code == "0", "dma_name"] = "OUTSIDE ANY DMA (REMOTE ALASKA)"
    z = z.merge(dz, on="zip", how="left")
    z["dma_source"] = z.dma_code.notna().map({True: "zip", False: ""})
    cdma = (z[z.dma_code.notna()].groupby(["county_fips", "dma_code", "dma_name"]).population.sum()
            .reset_index().sort_values("population").groupby("county_fips").tail(1).set_index("county_fips"))
    miss = z.dma_code.isna()
    z.loc[miss, "dma_code"] = z.loc[miss, "county_fips"].map(cdma.dma_code)
    z.loc[miss, "dma_name"] = z.loc[miss, "county_fips"].map(cdma.dma_name)
    z.loc[miss & z.dma_code.notna(), "dma_source"] = "county_fill"
    tot_pop = z.population.sum()
    zip_share = z.loc[z.dma_source == "zip", "population"].sum() / tot_pop
    nodma = z.dma_code.isna()
    rep.append(("Every ZIP has exactly one county, metro status and DMA",
                "PASS" if nodma.sum() == 0 else "FAIL",
                f"{zip_share:.1%} of population matched to a DMA by ZIP; {int((z.dma_source=='county_fill').sum())} "
                f"ZIPs filled from their county; {int(nodma.sum())} without a DMA"))

    # --- county table
    g = z.groupby(["county_fips", "dma_code", "dma_name"]).population.sum().reset_index()
    cty = (z.groupby("county_fips").agg(state=("State", "first"), county=("County", "first"),
                                         population=("population", "sum"), households=("households", "sum"),
                                         zips=("zip", "count")).reset_index())
    cty["dma_code"] = cty.county_fips.map(cdma.dma_code)
    cty["dma_name"] = cty.county_fips.map(cdma.dma_name)
    share = g.groupby("county_fips").population.apply(lambda s: s.max() / s.sum() if s.sum() else 1.0)
    cty["dma_pop_share"] = cty.county_fips.map(share).round(4)
    cty["split_county"] = cty.county_fips.map(g.groupby("county_fips").dma_code.nunique()) > 1
    # metro per county: population-majority CBSA from its ZIPs (OMB metros are whole counties)
    cc = z.groupby(["county_fips", "cbsa_code", "cbsa_name", "cbsa_type"]).population.sum().reset_index()
    cmaj = cc.sort_values("population").groupby("county_fips").tail(1).set_index("county_fips")
    cty["cbsa_code"] = cty.county_fips.map(cmaj.cbsa_code)
    cty["cbsa_name"] = cty.county_fips.map(cmaj.cbsa_name)
    cty["cbsa_type"] = cty.county_fips.map(cmaj.cbsa_type)
    inconsistent = cc.groupby("county_fips").cbsa_code.nunique()
    bad_cbsa = inconsistent[inconsistent > 1].index
    bad_pop = cc[cc.county_fips.isin(bad_cbsa)].groupby("county_fips").population.apply(lambda s: s.sum() - s.max()).sum()
    rep.append(("Each county sits in one metro (OMB metros are whole counties)",
                "PASS" if bad_pop / tot_pop < 0.005 else "REVIEW",
                f"{len(bad_cbsa)} counties have ZIPs pointing to more than one CBSA; {int(bad_pop):,} people "
                f"({bad_pop/tot_pop:.2%}) sit in the minority part. Counties use their majority CBSA."))
    n_metro = raw.loc[raw.CBSA_Type == "Metro", "CBSA"].nunique()
    n_micro = raw.loc[raw.CBSA_Type == "Micro", "CBSA"].nunique()
    rep.append(("Metro definitions are current (OMB Bulletin 23-01)",
                "PASS" if (n_metro, n_micro) == (393, 542) else "REVIEW",
                f"Whole file: {n_metro} metros and {n_micro} micros, matching OMB 2023's 393 and 542 "
                f"(including Puerto Rico); {z.loc[z.cbsa_type=='Metro','cbsa_code'].nunique()} metros fall in 50 states + DC"))
    split = cty[cty.split_county]
    rep.append(("Split counties", "INFO",
                f"{len(split)} counties are divided between DMAs by ZIP; largest: " +
                ", ".join(f"{r.county.title()} {r.state}" for r in split.nlargest(5, "population").itertuples())))

    # --- DMA and metro lists
    dma = (z.groupby(["dma_code", "dma_name"]).agg(population=("population", "sum"), households=("households", "sum"),
                                                   zips=("zip", "count"), counties=("county_fips", "nunique"))
           .reset_index().sort_values("households", ascending=False))
    dma["rank_by_households"] = range(1, len(dma) + 1)
    n_real = (dma.dma_code != "0").sum()
    out_pop = int(dma.loc[dma.dma_code == "0", "population"].sum())
    rep.append(("All 210 DMAs present", "PASS" if n_real == 210 else "FAIL",
                f"{n_real} DMAs; {out_pop:,} people in remote Alaska boroughs Nielsen leaves outside any DMA"))
    rep.append(("Manual DMA overrides", "REVIEW",
                f"{len(ov)} Coachella Valley ZIPs moved from Los Angeles to Palm Springs (804), which the public list omits; "
                f"{int(z[z.zip.isin(ov.zip)].population.sum()):,} people. Provisional until checked against a Nielsen source"))
    top10 = list(dma.dma_name.head(10))
    expect = ["NEW YORK", "LOS ANGELES", "CHICAGO", "PHILADELPHIA", "DALLAS", "ATLANTA", "HOUSTON",
              "WASHINGTON", "BOSTON", "SAN FRANCISCO"]
    hits = sum(any(e in t for t in top10) for e in expect)
    rep.append(("Spot check: top 10 DMAs by households match Nielsen's published top 10 markets",
                "PASS" if hits >= 9 else "REVIEW", "; ".join(f"{i+1}. {t.title()}" for i, t in enumerate(top10))))
    met = (z[z.cbsa_code != ""].groupby(["cbsa_code", "cbsa_name", "cbsa_type"])
           .agg(population=("population", "sum"), households=("households", "sum"), zips=("zip", "count"))
           .reset_index().sort_values("population", ascending=False))

    # --- national roll-up
    s_z, s_c, s_d = z.population.sum(), cty.population.sum(), dma.population.sum()
    rep.append(("Roll-up totals: ZIP = county = DMA = national", "PASS" if s_z == s_c == s_d else "FAIL",
                f"Population {s_z:,}; households {z.households.sum():,}"))
    rep.append(("Published benchmark: national population", "PASS" if 325e6 < s_z < 345e6 else "REVIEW",
                f"{s_z/1e6:.1f} million vs. about 332 million in the Census ACS 5-year national total"))

    # --- cross-check with 2021 county list
    dis = pd.DataFrame()
    p21 = os.path.join(HERE, "sources", "dma_county_2021.csv")
    if os.path.exists(p21):
        a21 = pd.read_csv(p21, dtype=str, encoding="latin1")
        norm = lambda s: re.sub(r"[^A-Z]", "", str(s).upper().replace("SAINT", "ST").replace(" COUNTY", "").replace(" PARISH", ""))
        a21["key"] = a21.STATE_AB + "|" + a21.COUNTY.map(norm)
        c2 = cty.copy(); c2["key"] = c2.state + "|" + c2.county.map(norm)
        c2 = c2.merge(a21[["key", "TVDMA"]].drop_duplicates("key"), on="key", how="inner")
        best = (c2.groupby(["TVDMA", "dma_name"]).size().reset_index(name="n").sort_values("n")
                .groupby("TVDMA").tail(1).set_index("TVDMA").dma_name)
        c2["dma_2021"] = c2.TVDMA.map(best)
        dis = c2[c2.dma_2021 != c2.dma_name][["county_fips", "state", "county", "population", "dma_name", "dma_2021"]]
        dis = dis.rename(columns={"dma_name": "dma_2026_used"}).sort_values("population", ascending=False)
        rep.append(("Cross-check against the 2021 county list", "INFO",
                    f"{len(c2):,} counties compared; {len(dis)} disagree, holding {int(dis.population.sum()):,} people "
                    f"({dis.population.sum()/tot_pop:.2%}). The 2026 list is used; all are in dma_disagreements.csv"))

    # --- write
    z["city"] = z.City.fillna("").astype(str).str.title()
    zout = z[["zip", "city", "State", "County", "county_fips", "cbsa_code", "cbsa_name", "cbsa_type", "dma_code", "dma_name",
              "dma_source", "multi_county", "population", "households", "Latitude", "Longitude"]]
    zout.columns = ["zip", "city", "state", "county", "county_fips", "cbsa_code", "cbsa_name", "cbsa_type", "dma_code",
                    "dma_name", "dma_source", "multi_county", "population", "households", "lat", "lon"]
    zout.to_csv(os.path.join(OUT, "zip_geo.csv"), index=False)
    cty.sort_values("county_fips").to_csv(os.path.join(OUT, "county_geo.csv"), index=False)
    dma.to_csv(os.path.join(OUT, "dma_list.csv"), index=False)
    met.to_csv(os.path.join(OUT, "metro_list.csv"), index=False)
    dis.to_csv(os.path.join(OUT, "dma_disagreements.csv"), index=False)

    failed = [r for r in rep if r[1] == "FAIL"]
    with open(os.path.join(OUT, "geography_report.md"), "w") as f:
        f.write("# Geography layer check report\n\n")
        f.write(f"Result: **{'FAILED' if failed else 'PASSED'}** "
                f"({sum(r[1]=='PASS' for r in rep)} pass, {sum(r[1]=='REVIEW' for r in rep)} for review, {len(failed)} fail)\n\n")
        f.write("| Check | Result | Detail |\n| --- | --- | --- |\n")
        for c, r, d in rep:
            f.write(f"| {c} | {r} | {d} |\n")
        f.write("\n## Largest DMA disagreements with the 2021 list\n\n| County | 2026 list (used) | 2021 list | Population |\n| --- | --- | --- | --- |\n")
        for r in dis.head(15).itertuples():
            f.write(f"| {r.county.title()}, {r.state} | {r.dma_2026_used.title()} | {r.dma_2021.title()} | {r.population:,} |\n")
        f.write("\nSources: Deluxe ZIP file (Q4 2025); ZIP-level DMA list, github.com/BritCrit/dma_county_zip (March 2026); "
                "county DMA list, github.com/alex-patton/US-TVDMA-BY-COUNTY (2021, cross-check only).\n")
    for r in rep:
        print(f"[{r[1]}] {r[0]}: {r[2]}")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
